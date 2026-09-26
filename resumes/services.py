"""
Resume text extraction service.

Kept separate from views/models so the extraction logic can be reused by
the Phase 8 AI analyzer and unit-tested in isolation.
"""
import io
from pypdf import PdfReader
from docx import Document


class ResumeExtractionError(Exception):
    pass


def extract_text_from_pdf(file_obj):
    """Extract text from a PDF file object (Django FieldFile or BytesIO)."""
    try:
        file_obj.seek(0)
        reader = PdfReader(file_obj)
        text_parts = []
        for page in reader.pages:
            page_text = page.extract_text()
            if page_text:
                text_parts.append(page_text)
        text = "\n".join(text_parts).strip()
        if not text:
            raise ResumeExtractionError(
                "No extractable text found. The PDF may be a scanned image without a text layer."
            )
        return text
    except ResumeExtractionError:
        raise
    except Exception as exc:
        raise ResumeExtractionError(f"Failed to read PDF: {exc}")


def extract_text_from_docx(file_obj):
    """Extract text from a DOCX file object (Django FieldFile or BytesIO)."""
    try:
        file_obj.seek(0)
        document = Document(file_obj)
        paragraphs = [p.text for p in document.paragraphs if p.text.strip()]

        # Also pull text from tables, since resumes often use table layouts for skills/experience.
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        paragraphs.append(cell.text.strip())

        text = "\n".join(paragraphs).strip()
        if not text:
            raise ResumeExtractionError("No extractable text found in the DOCX file.")
        return text
    except ResumeExtractionError:
        raise
    except Exception as exc:
        raise ResumeExtractionError(f"Failed to read DOCX: {exc}")


def extract_text(file_obj, file_type):
    """Dispatch to the correct extractor based on file_type ('pdf' or 'docx')."""
    file_type = file_type.lower().lstrip('.')
    if file_type == 'pdf':
        return extract_text_from_pdf(file_obj)
    elif file_type == 'docx':
        return extract_text_from_docx(file_obj)
    else:
        raise ResumeExtractionError(f"Unsupported file type: {file_type}")


def clean_extracted_text(text):
    """Basic normalization: collapse excessive blank lines/whitespace."""
    lines = [line.strip() for line in text.splitlines()]
    lines = [line for line in lines if line]
    return "\n".join(lines)
